#!/usr/bin/env python3
"""Check an HCIG email or report against the house voice rules.

usage: check_writing.py FILE [--partner] [--one-page]

  --partner   also fail on any emoji (partner, insurer, hotel, clinic, guest)
  --one-page  warn when the text is likely longer than one A4 page

Reads .md, .txt, .html and .docx. Exit code 1 when a hard rule fails.
"""
import re
import sys
from pathlib import Path

BANNED = [
    "i hope this email finds you well", "i hope you are well", "just wanted to",
    "world-class", "world class", "cutting-edge", "state-of-the-art", "seamless",
    "leverage", "synergy", "delve", "in today's", "it is important to note",
    "please do not hesitate", "feel free to", "accredited by", "certified by",
]
EMOJI = re.compile("[\U0001F300-\U0001FAFF☀-➿]")
DASH = re.compile("[—–]")


def read(path: Path) -> str:
    if path.suffix.lower() == ".docx":
        import docx  # python-docx
        d = docx.Document(str(path))
        parts = [p.text for p in d.paragraphs]
        for t in d.tables:
            for row in t.rows:
                parts.append(" | ".join(c.text for c in row.cells))
        return "\n".join(parts)
    text = path.read_text(encoding="utf-8", errors="replace")
    if path.suffix.lower() in (".html", ".htm"):
        text = re.sub(r"<(script|style)[^>]*>.*?</\1>", " ", text, flags=re.S | re.I)
        text = re.sub(r"<br\s*/?>|</(p|div|li|h\d|tr)>", "\n", text, flags=re.I)
        text = re.sub(r"<[^>]+>", " ", text)
    return text


def prose_lines(text: str):
    """Lines that are prose: skip code fences, tables, headings, links-only."""
    in_code = False
    for n, line in enumerate(text.splitlines(), 1):
        s = line.strip()
        if s.startswith("```"):
            in_code = not in_code
            continue
        if in_code or not s or s.startswith(("|", "#", "http", "---")):
            continue
        yield n, s


def main() -> int:
    # Windows consoles default to cp1252 and crash on emoji or dashes
    for stream in (sys.stdout, sys.stderr):
        try:
            stream.reconfigure(encoding="utf-8", errors="replace")
        except AttributeError:
            pass
    args = sys.argv[1:]
    if not args or args[0].startswith("-"):
        print(__doc__)
        return 2
    path = Path(args[0])
    partner = "--partner" in args
    one_page = "--one-page" in args
    text = read(path)
    fails, warns = [], []

    for n, line in enumerate(text.splitlines(), 1):
        if DASH.search(line):
            fails.append(f"line {n}: em or en dash: {line.strip()[:80]}")
        if partner and EMOJI.search(line):
            fails.append(f"line {n}: emoji in partner-facing text")
        low = line.lower()
        for b in BANNED:
            if b in low:
                (fails if b in ("accredited by", "certified by") else warns).append(
                    f"line {n}: '{b}'")

    para = []
    for n, s in prose_lines(text + "\n\n"):
        for sentence in re.split(r"(?<=[.!?])\s+", re.sub(r"^[-*\d.)\s]+", "", s)):
            words = len(sentence.split())
            if words > 25:
                warns.append(f"line {n}: sentence of {words} words, split it")
    block = []
    for line in text.splitlines() + [""]:
        s = line.strip()
        if s and not s.startswith(("|", "#", "-", "*", "```")) and not re.match(r"^\d+\.", s):
            block.append(s)
        else:
            if sum(len(x) for x in block) > 320:
                warns.append(f"paragraph of {sum(len(x) for x in block)} characters: headline plus one note")
            block = []

    words = len(re.findall(r"\w+", text))
    if one_page and words > 450:
        warns.append(f"{words} words: likely more than one page, cut")

    if len(EMOJI.findall(text)) > 1 and not partner:
        warns.append("more than one emoji")

    for f in fails:
        print("FAIL ", f)
    for w in warns:
        print("warn ", w)
    print(f"\n{path.name}: {words} words · {len(fails)} fail · {len(warns)} warn")
    return 1 if fails else 0


if __name__ == "__main__":
    sys.exit(main())
